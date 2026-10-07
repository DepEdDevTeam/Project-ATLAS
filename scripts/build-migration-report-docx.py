"""Convert a completed migration report to an editable Word document."""
import re
import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT


def inline(paragraph, text):
    pattern = r'(\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*|`[^`]+`)'
    for token in re.split(pattern, text):
        if not token:
            continue
        match = re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)', token)
        if match:
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), paragraph.part.relate_to(match[2], RT.HYPERLINK, is_external=True))
            run = OxmlElement('w:r')
            props = OxmlElement('w:rPr')
            color = OxmlElement('w:color')
            color.set(qn('w:val'), '195D86')
            props.append(color)
            underline = OxmlElement('w:u')
            underline.set(qn('w:val'), 'single')
            props.append(underline)
            run.append(props)
            content = OxmlElement('w:t')
            content.text = match[1]
            run.append(content)
            link.append(run)
            paragraph._p.append(link)
        else:
            run = paragraph.add_run(token.strip('*') if token.startswith('**') else token.strip('`'))
            if token.startswith('**'):
                run.bold = True


def table(doc, lines):
    rows = [[cell.strip() for cell in line.strip().strip('|').split('|')] for line in lines]
    rows = [r for r in rows if not all(re.fullmatch(r':?-+:?', cell) for cell in r)]
    result = doc.add_table(rows=0, cols=len(rows[0]))
    result.autofit = False
    widths = [1.25, 1.70, 1.70, 1.70] if len(rows[0]) == 4 else [1.15, .70, .70, 1.05, 1.05, 1.70]
    for col, width in zip(result.columns, widths):
        col.width = Inches(width)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement(f'w:{edge}')
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            border.set(qn(f'w:{key}'), value)
        borders.append(border)
    result._tbl.tblPr.append(borders)
    for index, values in enumerate(rows):
        row = result.add_row()
        if index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge in ('top', 'bottom', 'left', 'right'):
                node = OxmlElement(f'w:{edge}')
                node.set(qn('w:w'), '100')
                node.set(qn('w:type'), 'dxa')
                margins.append(node)
            props.append(margins)
            fill = OxmlElement('w:shd')
            fill.set(qn('w:fill'), '233E50' if index == 0 else ('F1F5F7' if index % 2 else 'FFFFFF'))
            props.append(fill)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(value)
            r.font.size = Pt(10)
            if index == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def main():
    source = Path(sys.argv[1]).resolve()
    destination = Path(sys.argv[2]).resolve()
    text = source.read_text(encoding='utf-8')
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.7)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'List Bullet'):
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    doc.styles['Title'].font.size = Pt(23)
    for name, size in [('Heading 1', 15), ('Heading 2', 12)]:
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(7)
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith('|'):
            block = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                block.append(lines[i])
                i += 1
            table(doc, block)
            continue
        if line.startswith('# '):
            doc.add_paragraph('ATLAS learner migration scenario for Region IV A from 2027 to 2035', 'Title')
            doc.add_paragraph('This report examines CALABARZON learner enrollment under assumed net migration rates, using DepEd aggregates and five research studies. The results are conditional what-if scenarios, not a calibrated migration forecast.')
        elif line.startswith('##'):
            level = 2 if line.startswith('###') else 1
            heading = re.sub(r'[^\w\s]', ' ', line.lstrip('#').strip())
            doc.add_heading(re.sub(r'\s+', ' ', heading), level)
        else:
            bullet = line.startswith('- ')
            p = doc.add_paragraph(style='List Bullet' if bullet else 'Normal')
            if bullet:
                p.paragraph_format.space_after = Pt(6)
            inline(p, line[2:] if bullet else line)
        i += 1
    doc.core_properties.title = 'ATLAS learner migration scenario for Region IV A from 2027 to 2035'
    doc.core_properties.subject = 'Multiple study migration sensitivity report'
    doc.core_properties.author = 'Project ATLAS'
    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)
    assert len(doc.tables) == 2
    print(f'Created {destination}')


if __name__ == '__main__':
    main()
